"""
Pattern 03: Live Automated Selenium Browser Login Runner
Executes an actual automated browser session using Selenium WebDriver on Jmix Bookstore,
consuming test accounts dynamically loaded from an Excel workbook (test_users.xlsx).

Toggle HEADLESS=0 in .env or your environment to view the live Chrome window.
"""

from __future__ import annotations
import os
from pathlib import Path
import secrets
import sys

from dotenv import load_dotenv
from openpyxl import load_workbook
from selenium import webdriver
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

# Ensure UTF-8 output on Windows consoles
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent.parent
CREDENTIALS_FILE = BASE_DIR / "data" / "test_users.xlsx"
LOGIN_URL = os.getenv("BOOKSTORE_LOGIN_URL", "https://demo.jmix.io/bookstore/login")
WAIT_SECONDS = int(os.getenv("SELENIUM_WAIT_SECONDS", "15"))

if not CREDENTIALS_FILE.is_file():
    raise FileNotFoundError(f"Credentials workbook not found: {CREDENTIALS_FILE}")


def get_random_credentials() -> tuple[str, str]:
    """Read a random non-empty User/Password row from the workbook."""
    workbook = load_workbook(CREDENTIALS_FILE, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        if headers[:2] != ("User", "Password"):
            raise ValueError(f"Active sheet columns must be 'User' and 'Password', found {headers!r}")
        rows = [
            (str(user).strip(), str(password).strip())
            for user, password, *_ in sheet.iter_rows(min_row=2, values_only=True)
            if user is not None and password is not None
            and str(user).strip() and str(password).strip()
        ]
    finally:
        workbook.close()

    if not rows:
        raise ValueError(f"No complete credential rows found in {CREDENTIALS_FILE}")
    return secrets.choice(rows)


def build_driver() -> webdriver.Chrome:
    """Configures and initializes Chrome WebDriver."""
    options = webdriver.ChromeOptions()
    headless_setting = os.getenv("HEADLESS", "1").strip().lower()
    if headless_setting not in {"0", "false", "no"}:
        options.add_argument("--headless=new")

    options.add_argument("--window-size=1440,1000")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    return webdriver.Chrome(options=options)


def run_login_check() -> None:
    username, password = get_random_credentials()
    print(f"Loaded credentials for user: '{username}' from {CREDENTIALS_FILE.name}")
    print(f"Launching Chrome (Headless: {os.getenv('HEADLESS', '1')})...")

    driver = build_driver()
    try:
        wait = WebDriverWait(driver, WAIT_SECONDS)
        print(f"Navigating to {LOGIN_URL}...")
        driver.get(LOGIN_URL)

        # Locate form inputs
        username_input = wait.until(EC.visibility_of_element_located((By.ID, "username")))
        password_input = wait.until(EC.visibility_of_element_located((By.ID, "password")))

        username_input.clear()
        username_input.send_keys(username)
        password_input.clear()
        password_input.send_keys(password)

        submit = wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, "button[type='submit']")))
        print("Submitting login form...")
        submit.click()

        # Verify login completion
        try:
            wait.until(
                lambda browser: "/login" not in browser.current_url.lower()
                or not browser.find_elements(By.ID, "password")
            )
        except TimeoutException as exc:
            raise AssertionError(
                "Login did not complete before timeout. Check form locators and credentials."
            ) from exc

        print(f"✅ Login check passed! Current URL: {driver.current_url}")
    finally:
        driver.quit()
        print("Browser session closed.")


if __name__ == "__main__":
    run_login_check()
